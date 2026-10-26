import os
import time
import uuid
from pathlib import Path

import requests
from openpyxl import load_workbook, Workbook

BASE_URL = os.getenv("CHATBOT_URL", "http://127.0.0.1:8000").rstrip("/")
INPUT_FILE = Path(os.getenv("CHATBOT_INPUT", "chatbot_200_test_questions.xlsx"))
OUTPUT_FILE = Path(os.getenv("CHATBOT_OUTPUT", "chatbot_200_test_results.xlsx"))
TIMEOUT = 120


def load_questions():
    wb = load_workbook(INPUT_FILE, read_only=True, data_only=True)
    ws = wb["Questions"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()
    return [
        {"No": r[0], "Category": r[1] or "", "Question": str(r[2]).strip()}
        for r in rows if r[2]
    ]


def ask(question, session_id):
    started = time.perf_counter()
    try:
        response = requests.post(
            f"{BASE_URL}/chat",
            json={"session_id": session_id, "question": question, "k": 10},
            timeout=TIMEOUT,
        )
        elapsed = round(time.perf_counter() - started, 3)

        try:
            data = response.json()
        except ValueError:
            data = {}

        if response.ok and isinstance(data, dict):
            return {
                "HTTP Status": response.status_code,
                "Answer": data.get("answer", data.get("response", "")),
                "Source": data.get("source", ""),
                "Sources": str(data.get("sources", [])),
                "Time (sec)": elapsed,
                "Error": "",
            }

        return {
            "HTTP Status": response.status_code,
            "Answer": "",
            "Source": "",
            "Sources": "",
            "Time (sec)": elapsed,
            "Error": str(data.get("detail", response.text[:1000])),
        }

    except Exception as exc:
        return {
            "HTTP Status": "",
            "Answer": "",
            "Source": "",
            "Sources": "",
            "Time (sec)": round(time.perf_counter() - started, 3),
            "Error": str(exc),
        }


def main():
    questions = load_questions()
    session_id = str(uuid.uuid4())

    print("=" * 70)
    print("200 QUESTION CHATBOT TEST")
    print(f"API: {BASE_URL}/chat")
    print(f"Session: {session_id}")
    print(f"Questions: {len(questions)}")
    print("=" * 70)

    results = []
    started = time.perf_counter()

    for index, item in enumerate(questions, 1):
        print(f"[{index}/{len(questions)}] {item['Question']}")

        result = ask(item["Question"], session_id)

        results.append({
            "No": item["No"],
            "Category": item["Category"],
            "Question": item["Question"],
            **result,
        })

        if result["Error"]:
            print("  ERROR:", result["Error"])
        else:
            print("  ANSWER:", str(result["Answer"]).replace("\n", " ")[:180])
            print("  SOURCE:", result["Source"])
            print("  TIME:", result["Time (sec)"], "sec")

    total_time = round(time.perf_counter() - started, 3)

    out = Workbook()
    ws = out.active
    ws.title = "Results"

    headers = [
        "No", "Category", "Question", "Answer", "Source",
        "Sources", "HTTP Status", "Time (sec)", "Error"
    ]
    ws.append(headers)

    for row in results:
        ws.append([row.get(h, "") for h in headers])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    widths = [8, 26, 65, 90, 22, 45, 14, 14, 45]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + idx)].width = width

    summary = out.create_sheet("Summary")
    successful = sum(1 for r in results if r["HTTP Status"] == 200 and not r["Error"])
    errors = len(results) - successful
    duckdb = sum(1 for r in results if "duckdb" in str(r["Source"]).lower())
    rag = sum(
        1 for r in results
        if "knowledge" in str(r["Source"]).lower() or "rag" in str(r["Source"]).lower()
    )

    summary.append(["Metric", "Value"])
    summary.append(["Total questions", len(results)])
    summary.append(["Successful responses", successful])
    summary.append(["Errors", errors])
    summary.append(["DuckDB responses", duckdb])
    summary.append(["RAG/Knowledge responses", rag])
    summary.append(["Total time (sec)", total_time])

    out.save(OUTPUT_FILE)

    print("=" * 70)
    print("DONE")
    print("Output:", OUTPUT_FILE.resolve())
    print("Successful:", successful)
    print("Errors:", errors)
    print("Total time:", total_time, "sec")
    print("=" * 70)


if __name__ == "__main__":
    main()
