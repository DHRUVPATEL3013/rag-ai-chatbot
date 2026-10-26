from pathlib import Path
import time
import uuid

import requests
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment


API_URL = "http://127.0.0.1:8000/chat"

# scripts/ -> project root
BASE_DIR = Path(__file__).resolve().parent.parent

QUESTIONS_FILE = (
    BASE_DIR / "ai_guidance_200_alternating_rag_duckdb_benchmark.xlsx"
)

RESULT_FILE = (
    BASE_DIR / "ai_guidance_200_alternating_rag_duckdb_results.xlsx"
)

TIMEOUT = 120
DELAY = 0.25


def main():

    if not QUESTIONS_FILE.exists():
        raise FileNotFoundError(
            f"Missing benchmark file:\n{QUESTIONS_FILE}"
        )

    wb = load_workbook(
        QUESTIONS_FILE,
        read_only=True
    )

    ws = wb["Questions"]

    session_id = str(uuid.uuid4())

    results = []

    print("=" * 75)
    print("AI GUIDANCE CHATBOT - 200 QUESTION TEST")
    print(f"API: {API_URL}")
    print(f"Questions: {QUESTIONS_FILE}")
    print(f"Output: {RESULT_FILE}")
    print(f"Session: {session_id}")
    print("=" * 75)

    for row in ws.iter_rows(
        min_row=2,
        values_only=True
    ):

        # New benchmark format:
        # No | Route | Question | Test Intent

        no, expected_route, question, test_intent = row

        started = time.perf_counter()

        try:

            response = requests.post(
                API_URL,
                json={
                    "question": question,
                    "session_id": session_id
                },
                timeout=TIMEOUT
            )

            elapsed = round(
                time.perf_counter() - started,
                3
            )

            try:
                data = response.json()
            except ValueError:
                data = {}

            answer = str(
                data.get("answer") or ""
            )

            source = str(
                data.get("source") or ""
            )

            detail = data.get(
                "detail",
                ""
            )

            status = "PASS"
            reason = ""

            # ==========================================
            # HTTP FAILURE
            # ==========================================

            if response.status_code != 200:

                status = "FAIL"

                reason = (
                    f"HTTP {response.status_code}: "
                    f"{detail or data}"
                )

            # ==========================================
            # EMPTY ANSWER
            # ==========================================

            elif not answer.strip():

                status = "FAIL"

                reason = "Empty answer"

            # ==========================================
            # DUCKDB VALIDATION
            # ==========================================

            elif expected_route == "DB":

                route_text = (
                    f"{source} {answer}"
                ).lower()

                # Your API may return a source such as
                # DuckDB, Excel, Structured, etc.

                db_indicators = [
                    "duckdb",
                    "excel",
                    "structured"
                ]

                if not any(
                    indicator in route_text
                    for indicator in db_indicators
                ):

                    status = "FAIL"

                    reason = (
                        "Expected DuckDB/structured "
                        f"route but returned source={source!r}"
                    )

            # ==========================================
            # RAG VALIDATION
            # ==========================================

            elif expected_route == "RAG":

                route_text = (
                    f"{source} {answer}"
                ).lower()

                rag_indicators = [
                    "internal knowledge",
                    "document",
                    "rag",
                    "policy"
                ]

                if not any(
                    indicator in route_text
                    for indicator in rag_indicators
                ):

                    status = "CHECK"

                    reason = (
                        "Answer returned but RAG source "
                        f"could not be confirmed. "
                        f"source={source!r}"
                    )

            # ==========================================
            # PRINT RESULT
            # ==========================================

            print(
                f"{no:03} | "
                f"{status:5} | "
                f"HTTP {response.status_code:3} | "
                f"{elapsed:6}s | "
                f"{expected_route:3} | "
                f"{question}"
            )

            results.append(
                [
                    no,
                    expected_route,
                    question,
                    test_intent,
                    response.status_code,
                    status,
                    reason,
                    answer,
                    source,
                    elapsed,
                    data.get(
                        "session_id",
                        session_id
                    )
                ]
            )

        except requests.RequestException as exc:

            elapsed = round(
                time.perf_counter() - started,
                3
            )

            print(
                f"{no:03} | "
                f"FAIL  | "
                f"API ERROR | "
                f"{question} | "
                f"{exc}"
            )

            results.append(
                [
                    no,
                    expected_route,
                    question,
                    test_intent,
                    "",
                    "FAIL",
                    str(exc),
                    "",
                    "",
                    elapsed,
                    session_id
                ]
            )

        time.sleep(DELAY)

    # ==============================================
    # CREATE RESULT WORKBOOK
    # ==============================================

    out = load_workbook(
        QUESTIONS_FILE
    )

    if "Results" in out.sheetnames:
        del out["Results"]

    if "Run Summary" in out.sheetnames:
        del out["Run Summary"]

    # ==============================================
    # RESULTS SHEET
    # ==============================================

    rs = out.create_sheet(
        "Results"
    )

    headers = [
        "No",
        "Expected Route",
        "Question",
        "Test Intent",
        "HTTP Status",
        "Status",
        "Reason",
        "Answer",
        "Returned Source",
        "Latency Seconds",
        "Session ID"
    ]

    rs.append(headers)

    for item in results:
        rs.append(item)

    # Header formatting

    for cell in rs[1]:

        cell.font = Font(
            bold=True
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="D9EAF7"
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    rs.freeze_panes = "A2"

    rs.auto_filter.ref = rs.dimensions

    widths = {
        "A": 7,
        "B": 18,
        "C": 90,
        "D": 40,
        "E": 12,
        "F": 12,
        "G": 60,
        "H": 90,
        "I": 25,
        "J": 18,
        "K": 40
    }

    for col, width in widths.items():
        rs.column_dimensions[col].width = width

    # ==============================================
    # SUMMARY
    # ==============================================

    summary = out.create_sheet(
        "Run Summary",
        0
    )

    total = len(results)

    passed = sum(
        x[5] == "PASS"
        for x in results
    )

    failed = sum(
        x[5] == "FAIL"
        for x in results
    )

    checks = sum(
        x[5] == "CHECK"
        for x in results
    )

    rag_total = sum(
        x[1] == "RAG"
        for x in results
    )

    db_total = sum(
        x[1] == "DB"
        for x in results
    )

    rag_pass = sum(
        x[1] == "RAG"
        and x[5] == "PASS"
        for x in results
    )

    db_pass = sum(
        x[1] == "DB"
        and x[5] == "PASS"
        for x in results
    )

    summary.append(
        ["Metric", "Value"]
    )

    summary.append(
        ["Total questions", total]
    )

    summary.append(
        ["RAG questions", rag_total]
    )

    summary.append(
        ["DuckDB questions", db_total]
    )

    summary.append(
        ["PASS", passed]
    )

    summary.append(
        ["FAIL", failed]
    )

    summary.append(
        ["CHECK", checks]
    )

    summary.append(
        ["RAG PASS", rag_pass]
    )

    summary.append(
        ["DuckDB PASS", db_pass]
    )

    summary.append(
        [
            "Pass rate excluding CHECK",
            round(
                100 * passed /
                max(1, passed + failed),
                2
            )
        ]
    )

    for cell in summary[1]:

        cell.font = Font(
            bold=True
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="D9EAF7"
        )

    summary.column_dimensions["A"].width = 40
    summary.column_dimensions["B"].width = 20

    # ==============================================
    # SAVE
    # ==============================================

    out.save(
        RESULT_FILE
    )

    print("=" * 75)
    print(
        f"Completed: {total}"
    )
    print(
        f"PASS: {passed}"
    )
    print(
        f"FAIL: {failed}"
    )
    print(
        f"CHECK: {checks}"
    )
    print(
        f"RAG PASS: {rag_pass}/{rag_total}"
    )
    print(
        f"DB PASS: {db_pass}/{db_total}"
    )
    print(
        f"Results saved to:\n{RESULT_FILE}"
    )
    print("=" * 75)


if __name__ == "__main__":
    main()